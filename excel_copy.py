#!/usr/bin/env python3
"""Copy an Excel workbook to a new file while preserving row and column
layout for every sheet.

Usage:
    python excel_copy.py source.xlsx destination.xlsx

The script reads each sheet from the source workbook and writes it to a new
workbook at the destination path, keeping all cell values in the same row and
column positions.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook, Workbook


def copy_workbook(src: Path, dest: Path) -> None:
    """Copy every sheet from *src* workbook to *dest* workbook.

    Parameters
    ----------
    src: Path
        Path to the existing Excel file to copy from.
    dest: Path
        Destination path for the new Excel file.
    """
    wb_src = load_workbook(filename=src, data_only=True)
    wb_dest = Workbook()

    # Remove the automatically created empty sheet in the new workbook
    default_sheet = wb_dest.active
    wb_dest.remove(default_sheet)

    for ws_src in wb_src.worksheets:
        ws_dest = wb_dest.create_sheet(title=ws_src.title)
        for row in ws_src.iter_rows(values_only=True):
            ws_dest.append(list(row))

    wb_dest.save(dest)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Copy all sheets from an Excel workbook to a new file, preserving row and column alignment."
    )
    parser.add_argument("source", type=Path, help="Path to the Excel file to copy")
    parser.add_argument(
        "destination", type=Path, help="Path where the new Excel file will be written"
    )
    args = parser.parse_args()

    copy_workbook(args.source, args.destination)


if __name__ == "__main__":
    main()
